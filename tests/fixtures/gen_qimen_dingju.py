import openpyxl, re, json

SRC = "/Users/goodcamper/Desktop/seishi-app/docs/source/奇門遁甲万年暦.xlsx"
OUT = "/Users/goodcamper/Desktop/seishi-app/tests/fixtures/qimen_dingju_manreki.json"

wb = openpyxl.load_workbook(SRC, data_only=True)
STEMS = "甲乙丙丁戊己庚辛壬癸"
BRANCHES = "子丑寅卯辰巳午未申酉戌亥"

FULLWIDTH_DIGITS = {"０":"0","１":"1","２":"2","３":"3","４":"4","５":"5","６":"6","７":"7","８":"8","９":"9"}

def clean(s):
    if s is None:
        return None
    return str(s).replace("　", "").replace(" ", "").strip()

def normalize_label(lbl):
    if lbl is None:
        return None
    out = "".join(FULLWIDTH_DIGITS.get(c, c) for c in lbl)
    out = out.replace("（閏）", "").replace("(閏)", "").strip()
    m = re.match(r"^([陽陰])(\d+)$", out)
    if not m:
        return None
    return {"dun": "陽遁" if m.group(1) == "陽" else "陰遁", "ju": int(m.group(2))}

sheet_years = {
    "2018年　平成30年　戊戌": 2018,
    "2019年　平成31年　己亥": 2019,
    "2020年　平成32年　庚子": 2020,
}
MONTH_COL_PAIRS = [(1, "jan_next"), (3, None), (5, None), (7, None), (9, None), (11, None),
                    (15, None), (17, None), (19, None), (21, None), (23, None), (25, None)]

def ganzhi_index(gz):
    si, bi = STEMS.index(gz[0]), BRANCHES.index(gz[1])
    for i in range(60):
        if i % 10 == si and i % 12 == bi:
            return i
    raise ValueError(gz)

rows = []
skipped_bad_ganzhi = []

# ---- 干支の連続性チェック ----------------------------------------------
# 干支は暦日ごとに必ず+1ずつ進む（例外なし）。この不変条件を使って、
# 「幹支の組み合わせとしては有効（STEMS/BRANCHESに含まれ、陰陽の対応も正しい）
# だが、前後の日と矛盾する」ような単純な転記ミス（1文字だけ違う干支など）を、
# ロジック側の判定を一切参照せずに機械的に検出する。
all_ganzhi_cells = []  # (date, ganzhi) すべて（ラベルの有無に関わらず全セル）
for sheetname, year in sheet_years.items():
    ws = wb[sheetname]
    month_header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    for col, tag in MONTH_COL_PAIRS:
        mheader = clean(month_header[col - 1])
        m = re.search(r"(\d+)月", mheader)
        month_num = int(m.group(1)) if m else None
        cal_year = year + 1 if tag == "jan_next" else year
        for r in range(5, ws.max_row + 1):
            ganzhi = ws.cell(row=r, column=col + 1).value
            day_of_month = r - 4
            if ganzhi is None:
                continue
            gz = clean(ganzhi)
            if len(gz) != 2 or gz[0] not in STEMS or gz[1] not in BRANCHES:
                continue
            si, bi = STEMS.index(gz[0]), BRANCHES.index(gz[1])
            if (si % 2) != (bi % 2):
                continue
            try:
                import datetime
                d = datetime.date(cal_year, month_num, day_of_month)
            except ValueError:
                continue
            all_ganzhi_cells.append((d, gz))

all_ganzhi_cells.sort(key=lambda x: x[0])
dedup = {}
for d, gz in all_ganzhi_cells:
    dedup.setdefault(d, gz)
chain = sorted(dedup.items())

break_pairs = []  # (prev_date, date) の組。この間で干支の連続性が破れている。
prev = None
for d, gz in chain:
    idx = ganzhi_index(gz)
    if prev is not None:
        pd, pidx = prev
        expected = (pidx + (d - pd).days) % 60
        if expected != idx:
            break_pairs.append((pd, d))
    prev = (d, idx)

# 1セルだけが壊れている場合、その前後2箇所で連続性が破れる
# （前日→当該日、当該日→翌日 の2ペア）。2つのペアに共通する日付が
# 「壊れているセルそのもの」。単発の転記ミスをここで機械的に一意特定する。
chain_break_dates = set()
for i in range(len(break_pairs) - 1):
    _, d1_end = break_pairs[i]
    d2_start, _ = break_pairs[i + 1]
    if d1_end == d2_start:
        chain_break_dates.add(d1_end)

for sheetname, year in sheet_years.items():
    ws = wb[sheetname]
    month_header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

    for col, tag in MONTH_COL_PAIRS:
        mheader = clean(month_header[col - 1])
        m = re.search(r"(\d+)月", mheader)
        month_num = int(m.group(1)) if m else None
        cal_year = year + 1 if tag == "jan_next" else year

        for r in range(5, ws.max_row + 1):
            label = ws.cell(row=r, column=col).value
            ganzhi = ws.cell(row=r, column=col + 1).value
            day_of_month = r - 4
            if ganzhi is None:
                continue
            gz = clean(ganzhi)
            if len(gz) != 2 or gz[0] not in STEMS or gz[1] not in BRANCHES:
                skipped_bad_ganzhi.append({"year": cal_year, "month": month_num, "day": day_of_month, "raw": gz,
                                            "reason": "invalid stem/branch character"})
                continue
            si, bi = STEMS.index(gz[0]), BRANCHES.index(gz[1])
            if (si % 2) != (bi % 2):
                skipped_bad_ganzhi.append({"year": cal_year, "month": month_num, "day": day_of_month, "raw": gz,
                                            "reason": "invalid stem/branch parity"})
                continue
            try:
                import datetime
                cell_date = datetime.date(cal_year, month_num, day_of_month)
            except ValueError:
                cell_date = None
            if cell_date is not None and cell_date in chain_break_dates:
                skipped_bad_ganzhi.append({"year": cal_year, "month": month_num, "day": day_of_month, "raw": gz,
                                            "reason": "chain-consistency break: 前後の日から機械的に逆算した"
                                                      "正しい干支と一致しない（単純な転記ミスと判定）"})
                continue

            parsed = normalize_label(clean(label))
            if parsed is None:
                continue  # 局数ラベルが記載されている日だけを fixture 化する

            leap = "閏" in (clean(label) or "")
            rows.append({
                "year": cal_year, "month": month_num, "day": day_of_month,
                "dayStem": gz[0], "dayBranch": gz[1],
                "expectedDun": parsed["dun"], "expectedJu": parsed["ju"],
                "leap": leap,
            })

# 既知の Excel 転記ミスへの上書き（局数.pdf で二重に確認済みの正しい値を優先する）。
# 2018-11-23（己未）: Excel記載「陰3」→ 小雪・下元の正しい局数は「陰2」
#   （docs/source/局数.pdf の2つの独立引用が一致して「陰2」を示す。既にsrc側の
#   YIN_JU_TABLE["小雪"]=[5,8,2] も同じ値に修正済み。Excelの「陰3」の方を
#   転記ミスと判定し、fixtureはPDF側の正しい値に合わせる）。
OVERRIDES = {
    (2018, 11, 23): {
        "expectedDun": "陰遁", "expectedJu": 2,
        "note": "Excel原本は「陰3」と記載されているが、局数.pdfの独立引用2箇所（統宗の訂正記事・"
                "個別列挙リストの両方）が一致して「小雪下元=2」を示しており、Excel側の転記ミスと判定した。"
                "この1件だけはPDF確定値（陰2）を正としてテストする（Excelの陰3には合わせない）。",
    },
}

override_applied = []
for row in rows:
    key = (row["year"], row["month"], row["day"])
    if key in OVERRIDES:
        row.update(OVERRIDES[key])
        override_applied.append(key)

missing_overrides = [k for k in OVERRIDES if k not in override_applied]

fixture = {
    "_meta": {
        "source": "docs/source/奇門遁甲万年暦.xlsx（2018〜2020年、3シート）",
        "description": (
            "奇門遁甲万年暦.xlsx から機械的に転記した検証専用 fixture。仕様ではない。"
            "production コード（dingju.ts）は Excel を実行時に一切参照しない（Excel非依存）。"
            "このfixtureは resolveDingju() が返す dun/ju が、実際の万年暦の運用結果と"
            "日単位で一致することを固定する回帰テストの検証データ。"
        ),
        "excludedBadGanzhi": skipped_bad_ganzhi,
        "overrides": [
            {"date": f"{y}-{m:02d}-{d:02d}", **v} for (y, m, d), v in OVERRIDES.items()
        ],
        "generationScript": "tests/fixtures/gen_qimen_dingju.py",
        "count": len(rows),
    },
    "rows": rows,
}

json.dump(fixture, open(OUT, "w"), ensure_ascii=False, indent=1)
print("rows:", len(rows))
print("excluded bad ganzhi cells:", len(skipped_bad_ganzhi))
print("overrides applied:", override_applied)
if missing_overrides:
    print("WARNING: override target date not found in extracted rows:", missing_overrides)
